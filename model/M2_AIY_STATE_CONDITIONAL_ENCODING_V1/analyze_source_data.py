#!/usr/bin/env python3
"""Animal-level blocked-CV decoding of motor state from Ji et al. source traces."""
from __future__ import annotations

import argparse
import hashlib
import itertools
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import roc_auc_score
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parents[2]
EXP = "M2_AIY_STATE_CONDITIONAL_ENCODING_V1"
RAW = ROOT / "data/raw/celegans/ji_etal_2021_elife_68848_v3"
DEFAULT_OUT = ROOT / "data/results" / EXP
FOLDS = 10
PURGE_SAMPLES = 20  # 10 seconds at 2 Hz
C = 1.0


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def num(v):
    if v is None or (isinstance(v, str) and v.strip().lower() in {"nan", ""}):
        return np.nan
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan


def parse_grouped_sheet(path: Path, sheet_name: str, layouts: list[tuple[int, int, int, int]],
                        group: str) -> list[dict]:
    ws = load_workbook(path, read_only=True, data_only=True)[sheet_name]
    # Sources contain a title row and a header row; numerical samples start on row 3.
    all_rows = list(ws.iter_rows(min_row=3, values_only=True))
    animals = []
    for n, (start, state_col, aiy_col, ava_col) in enumerate(layouts, 1):
        vals = [[num(row[start]), num(row[start + state_col]),
                 num(row[start + aiy_col]), num(row[start + ava_col])] for row in all_rows]
        a = np.asarray(vals, dtype=float)
        valid = np.isfinite(a[:, 0]) & np.isfinite(a[:, 1])
        a = a[valid]
        if len(a) == 0 or np.unique(a[:, 1]).size < 2:
            continue
        animals.append({"group": group, "animal_id": f"{group}_{n}", "time": a[:, 0],
                        "state": a[:, 1], "aiy": a[:, 2], "ava": a[:, 3]})
    return animals


def load_animals():
    wt_path = RAW / "elife-68848-fig2-data1-v3.xlsx"
    rim_path = RAW / "elife-68848-fig5-data1-v3.xlsx"
    wt_layouts = [(0, 1, 2, 3), (9, 1, 2, 3), (18, 1, 2, 3),
                  (27, 1, 2, 3), (36, 2, 3, 4), (46, 1, 2, 3)]
    rim_layouts = [(0, 1, 2, 3), (8, 1, 2, 3), (16, 1, 2, 3),
                   (24, 1, 2, 3), (32, 1, 2, 3)]
    wt = parse_grouped_sheet(wt_path, "Fig 2HI_const temp", wt_layouts, "WT")
    rim = parse_grouped_sheet(rim_path, "Fig 5C,S1B", rim_layouts, "RIM_ABLATED")
    return wt + rim, [wt_path, rim_path]


def blocked_oof_auc(state, neural, include_state=True):
    state = np.asarray(state, dtype=float)
    neural = np.asarray(neural, dtype=float)
    # Keep only forward/reversal labels; preserve original row positions for contiguous folds.
    prior = np.r_[np.nan, state[:-1]]
    idx = np.flatnonzero(np.isin(state, (-1.0, 1.0)) & np.isfinite(neural) & np.isfinite(prior))
    y = (state[idx] == 1.0).astype(int)
    xparts = [prior[idx, None]] if include_state else []
    xparts.append(neural[idx, None])
    x = np.concatenate(xparts, axis=1)
    pred = np.full(len(idx), np.nan)
    nraw = len(state)
    fold_bounds = np.linspace(0, nraw, FOLDS + 1, dtype=int)
    for fold in range(FOLDS):
        lo, hi = fold_bounds[fold], fold_bounds[fold + 1]
        test = (idx >= lo) & (idx < hi)
        train = (idx < max(0, lo - PURGE_SAMPLES)) | (idx >= min(nraw, hi + PURGE_SAMPLES))
        train &= ~test
        if test.sum() < 2 or np.unique(y[train]).size < 2 or np.unique(y[test]).size < 2:
            # AUC is calculated after pooling OOF predictions; a test fold need not contain both labels.
            if test.sum() < 2 or np.unique(y[train]).size < 2:
                continue
        scaler = StandardScaler().fit(x[train])
        clf = LogisticRegression(C=C, solver="liblinear", random_state=0).fit(scaler.transform(x[train]), y[train])
        pred[test] = clf.predict_proba(scaler.transform(x[test]))[:, 1]
    ok = np.isfinite(pred)
    if ok.sum() < 50 or np.unique(y[ok]).size < 2:
        raise ValueError("Insufficient out-of-fold samples for AUC")
    return float(roc_auc_score(y[ok], pred[ok])), int(ok.sum())


def bootstrap_diff(a, b, n=50000, seed=20261001):
    rng = np.random.default_rng(seed)
    a, b = np.asarray(a), np.asarray(b)
    draws = rng.choice(a, (n, len(a)), replace=True).mean(axis=1) - rng.choice(b, (n, len(b)), replace=True).mean(axis=1)
    return [float(x) for x in np.quantile(draws, [0.025, 0.975])]


def exact_perm(a, b):
    vals = np.r_[a, b]
    observed = float(np.mean(a) - np.mean(b))
    diffs = []
    for chosen in itertools.combinations(range(len(vals)), len(a)):
        mask = np.zeros(len(vals), dtype=bool)
        mask[list(chosen)] = True
        diffs.append(vals[mask].mean() - vals[~mask].mean())
    diffs = np.asarray(diffs)
    return observed, float(np.mean(np.abs(diffs) >= abs(observed) - 1e-15)), int(len(diffs))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    out = args.output_dir
    out.mkdir(parents=True, exist_ok=True)
    animals, source_files = load_animals()
    source_records = [{"path": str(p.relative_to(ROOT)), "bytes": p.stat().st_size, "sha256": sha256(p)} for p in source_files]
    (out / "source_manifest.json").write_text(json.dumps({"experiment_id": EXP, "source_files": source_records}, indent=2) + "\n")
    results = []
    for a in animals:
        if len(a["time"]) < 300 or not np.all(np.diff(a["time"]) > 0):
            raise ValueError(f"Invalid time axis: {a['animal_id']}")
        dts = np.diff(a["time"])
        if not 0.45 <= np.median(dts) <= 0.55:
            raise ValueError(f"Unexpected sample interval {np.median(dts)}: {a['animal_id']}")
        row = {"group": a["group"], "animal_id": a["animal_id"], "n_rows": len(a["time"]),
               "n_forward": int(np.sum(a["state"] == 1)), "n_reverse": int(np.sum(a["state"] == -1)),
               "n_transition_or_unknown": int(np.sum(a["state"] == 0)),
               "dt_median_s": float(np.median(dts))}
        for neuron in ("aiy", "ava"):
            base_auc, n_base = blocked_oof_auc(a["state"], a[neuron], include_state=False)
            joint_auc, n_joint = blocked_oof_auc(a["state"], a[neuron], include_state=True)
            state_only_auc, n_state = blocked_oof_auc(a["state"], np.zeros_like(a["state"]), include_state=True)
            # For the behavior-only comparator, call decoder directly on previous state as its only feature.
            row[f"{neuron}_only_auc"] = base_auc
            row[f"{neuron}_plus_previous_state_auc"] = joint_auc
            row[f"{neuron}_increment_over_previous_state_auc"] = joint_auc - state_only_auc
            row[f"{neuron}_oof_samples"] = min(n_base, n_joint, n_state)
        results.append(row)
    df = pd.DataFrame(results)
    df.to_csv(out / "animal_results.csv", index=False, float_format="%.10g")
    summary = {"experiment_id": EXP, "classification": "retrospective_public_source_data_reanalysis",
               "primary_endpoint": "AUC(previous motor state + AIY) - AUC(previous motor state)",
               "animal_count": df.groupby("group").size().to_dict(), "animal_results": results,
               "group_summary": {}, "primary_group_contrast": {},
               "source_files": source_records,
               "analysis": {"folds": FOLDS, "purge_samples": PURGE_SAMPLES, "logistic_C": C,
                            "solver": "liblinear", "feature_scaling": "train-fold only", "unit": "animal"}}
    for metric in ("aiy_only_auc", "aiy_increment_over_previous_state_auc", "ava_only_auc", "ava_increment_over_previous_state_auc"):
        summary["group_summary"][metric] = {}
        for group, d in df.groupby("group"):
            vals = d[metric].to_numpy(float)
            summary["group_summary"][metric][group] = {"mean": float(vals.mean()), "animal_values": vals.tolist()}
    for neuron, metric in (("AIY", "aiy_increment_over_previous_state_auc"), ("AVA", "ava_increment_over_previous_state_auc")):
        wt = df.loc[df.group == "WT", metric].to_numpy(float)
        rim = df.loc[df.group == "RIM_ABLATED", metric].to_numpy(float)
        obs, p2, nperm = exact_perm(wt, rim)
        summary["primary_group_contrast"][neuron] = {"mean_WT_minus_RIM_ABLATED": obs,
            "animal_bootstrap_95pct_interval": bootstrap_diff(wt, rim),
            "exact_two_sided_label_permutation_p_descriptive": p2, "allocations": nperm}
    (out / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    manifest = {"experiment_id": EXP, "runner_sha256": sha256(Path(__file__)),
                "contract_sha256": sha256(ROOT / "summery" / EXP / "CONTRACT.md"),
                "outputs": {p.name: sha256(p) for p in (out / "animal_results.csv", out / "summary.json")}}
    (out / "run_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(json.dumps({"animals": summary["animal_count"], "group_summary": summary["group_summary"],
                      "primary_group_contrast": summary["primary_group_contrast"]}, indent=2))


if __name__ == "__main__":
    main()
