#!/usr/bin/env python3
"""Independent source-to-summary verification for M2 AIY decoding reanalysis."""
from __future__ import annotations

import hashlib
import itertools
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import roc_auc_score
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parents[2]
EXP = "M2_AIY_STATE_CONDITIONAL_ENCODING_V1"
OUT = ROOT / "data/results" / EXP
RAW = ROOT / "data/raw/celegans/ji_etal_2021_elife_68848_v3"


def digest(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def fnum(x):
    try:
        if x is None or (isinstance(x, str) and x.lower() == "nan"):
            return np.nan
        return float(x)
    except (ValueError, TypeError):
        return np.nan


def source_series(path, sheet, layouts, label):
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    out = []
    for k, (start, s_col, ai_col, av_col) in enumerate(layouts, 1):
        arr = np.array([[fnum(r[start]), fnum(r[start+s_col]), fnum(r[start+ai_col]), fnum(r[start+av_col])] for r in rows])
        arr = arr[np.isfinite(arr[:, 0]) & np.isfinite(arr[:, 1])]
        if len(arr) == 0 or np.unique(arr[:, 1]).size < 2:
            continue
        out.append((f"{label}_{k}", arr[:, 0], arr[:, 1], arr[:, 2], arr[:, 3]))
    return out


def score(state, signal, with_previous):
    previous = np.r_[np.nan, state[:-1]]
    ix = np.flatnonzero(np.isin(state, [-1., 1.]) & np.isfinite(signal) & np.isfinite(previous))
    target = (state[ix] == 1).astype(int)
    features = np.column_stack([previous[ix], signal[ix]]) if with_previous else signal[ix, None]
    oof = np.full(len(ix), np.nan)
    bounds = np.linspace(0, len(state), 11, dtype=int)
    for fold in range(10):
        lo, hi = bounds[fold:fold+2]
        held = (ix >= lo) & (ix < hi)
        keep = ((ix < max(0, lo-20)) | (ix >= min(len(state), hi+20))) & ~held
        if held.sum() < 2 or np.unique(target[keep]).size != 2:
            continue
        scaler = StandardScaler().fit(features[keep])
        fit = LogisticRegression(C=1.0, solver="liblinear", random_state=0)
        fit.fit(scaler.transform(features[keep]), target[keep])
        oof[held] = fit.predict_proba(scaler.transform(features[held]))[:, 1]
    ok = np.isfinite(oof)
    if ok.sum() < 50:
        raise AssertionError("Too few held-out predictions")
    return roc_auc_score(target[ok], oof[ok]), int(ok.sum())


def exact_difference(a, b):
    values = np.r_[a, b]
    obs = float(np.mean(a)-np.mean(b))
    diff = []
    for chosen in itertools.combinations(range(len(values)), len(a)):
        m = np.zeros(len(values), dtype=bool)
        m[list(chosen)] = True
        diff.append(values[m].mean()-values[~m].mean())
    diff = np.asarray(diff)
    return obs, float(np.mean(np.abs(diff) >= abs(obs)-1e-15)), len(diff)


def main():
    expected = {
        "WT": source_series(RAW/"elife-68848-fig2-data1-v3.xlsx", "Fig 2HI_const temp",
                            [(0,1,2,3),(9,1,2,3),(18,1,2,3),(27,1,2,3),(36,2,3,4),(46,1,2,3)], "WT"),
        "RIM_ABLATED": source_series(RAW/"elife-68848-fig5-data1-v3.xlsx", "Fig 5C,S1B",
                            [(0,1,2,3),(8,1,2,3),(16,1,2,3),(24,1,2,3),(32,1,2,3)], "RIM_ABLATED")}
    assert {k:len(v) for k,v in expected.items()} == {"WT":5,"RIM_ABLATED":5}
    rows=[]
    for group, series in expected.items():
        for animal_id, t, state, aiy, ava in series:
            assert np.all(np.diff(t) > 0) and .45 <= np.median(np.diff(t)) <= .55
            rec={"group":group,"animal_id":animal_id}
            for neuron, signal in (("aiy",aiy),("ava",ava)):
                one,_=score(state,signal,False)
                joint,n=score(state,signal,True)
                base,_=score(state,np.zeros_like(state),True)
                rec[f"{neuron}_only_auc"]=one
                rec[f"{neuron}_plus_previous_state_auc"]=joint
                rec[f"{neuron}_increment_over_previous_state_auc"]=joint-base
                rec[f"{neuron}_oof_samples"]=n
            rows.append(rec)
    got=pd.read_csv(OUT/"animal_results.csv")
    ref=pd.DataFrame(rows)
    assert len(got)==10 and set(got.animal_id)==set(ref.animal_id)
    numeric=[c for c in ref.columns if c not in {"group","animal_id"}]
    merged=got.merge(ref,on=["group","animal_id"],suffixes=("_out","_ref"),validate="one_to_one")
    for c in numeric:
        assert np.allclose(merged[f"{c}_out"],merged[f"{c}_ref"],rtol=0,atol=1e-10), c
    summary=json.loads((OUT/"summary.json").read_text())
    for n,metric in (("AIY","aiy_increment_over_previous_state_auc"),("AVA","ava_increment_over_previous_state_auc")):
        a=got.loc[got.group=="WT",metric].to_numpy()
        b=got.loc[got.group=="RIM_ABLATED",metric].to_numpy()
        obs,p,nperm=exact_difference(a,b)
        rec=summary["primary_group_contrast"][n]
        assert np.isclose(obs,rec["mean_WT_minus_RIM_ABLATED"],atol=1e-12)
        assert np.isclose(p,rec["exact_two_sided_label_permutation_p_descriptive"],atol=1e-12)
        assert nperm==rec["allocations"]==252
    manifest=json.loads((OUT/"run_manifest.json").read_text())
    assert manifest["experiment_id"]==EXP
    assert manifest["runner_sha256"]==digest(ROOT/"model"/EXP/"analyze_source_data.py")
    assert manifest["contract_sha256"]==digest(ROOT/"summery"/EXP/"CONTRACT.md")
    for name,h in manifest["outputs"].items():
        assert h==digest(OUT/name), name
    for src in summary["source_files"]:
        assert src["sha256"]==digest(ROOT/src["path"])
    report={"status":"PASS","source_animal_count":{"WT":5,"RIM_ABLATED":5},
            "row_count":int(len(got)),"independently_recomputed_model_scores":True,
            "exact_permutation_allocations":252,"all_hashes_verified":True}
    (OUT/"POSTRUN_VERIFICATION.json").write_text(json.dumps(report,indent=2)+"\n")
    print(json.dumps(report,indent=2))


if __name__=="__main__":
    main()
